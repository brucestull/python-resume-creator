#!/usr/bin/env python3
"""
resume_builder.py — Build an ATS-friendly Word (.docx) resume from a JSON file.

Design goals (why it looks the way it does):
  * Single column, top-to-bottom reading order. No tables anywhere.
    Applicant Tracking Systems (ATS) linearize a document before parsing, so a
    single flowing column reads correctly; multi-column tables get scrambled.
  * Standard section headings ("Summary", "Experience", "Education",
    "Certifications", "Skills"). Parsers look for these near-exact words to
    bucket your content.
  * Standard font (Calibri), real bold/size for hierarchy instead of graphics.
  * Skills rendered as comma-separated rows (packed to a readable line width),
    which parsers split cleanly on the commas.
  * Contact links (email, phone, website, LinkedIn, GitHub) are clickable.

Usage:
    python resume_builder.py --input sample_one_page.json --output resume.docx
    python resume_builder.py -i my_resume.json -o out.docx --skills-per-row 8

Dependencies:
    pip install python-docx      (or: uv add python-docx)
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches
from docx.text.run import Run
from docx.opc.constants import RELATIONSHIP_TYPE as RT


# --------------------------------------------------------------------------- #
# Style configuration — tweak these to change the look without touching logic. #
# --------------------------------------------------------------------------- #
@dataclass
class Style:
    font_name: str = "Calibri"

    # Point sizes
    name_size: int = 24
    title_size: int = 12          # the professional title under the name
    contact_size: int = 10
    section_size: int = 12        # section headers (EXPERIENCE, etc.)
    body_size: int = 10.5
    entry_title_size: int = 11    # "Position — Company" line

    # Colors (RGB)
    accent: RGBColor = RGBColor(0x1F, 0x3B, 0x57)   # dark slate blue
    text: RGBColor = RGBColor(0x1A, 0x1A, 0x1A)
    muted: RGBColor = RGBColor(0x55, 0x55, 0x55)

    # Page geometry (US Letter with tight-but-safe margins)
    page_width_in: float = 8.5
    page_height_in: float = 11.0
    margin_in: float = 0.6

    # Skills layout
    skills_per_row: int | None = None   # if set, force N skills per row
    skills_line_width: int = 90         # else greedy-pack up to ~this many chars

    @property
    def content_width_in(self) -> float:
        return self.page_width_in - 2 * self.margin_in


# --------------------------------------------------------------------------- #
# Low-level docx helpers                                                       #
# --------------------------------------------------------------------------- #
def _set_run(run, *, size=None, bold=False, italic=False, color=None, font=None):
    """Apply font properties to a run in one call."""
    if font:
        run.font.name = font
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _tight(paragraph, *, before=0, after=0, line=None):
    """Control paragraph spacing so the resume stays compact."""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    return paragraph


def _add_bottom_border(paragraph):
    """
    Draw a thin rule under a paragraph (used beneath section headers).
    This is a *paragraph border*, NOT a table, so it is ATS-safe.
    """
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")        # 6 eighths of a point = 0.75pt
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F3B57")
    borders.append(bottom)
    p_pr.append(borders)


def _right_tab(paragraph, content_width_in: float):
    """Register a right-aligned tab stop at the right margin."""
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(content_width_in), WD_TAB_ALIGNMENT.RIGHT
    )


def add_hyperlink(paragraph, text, url):
    """
    Return an ordinary, styleable Run that is wrapped in a hyperlink to *url*.

    python-docx has no high-level hyperlink API. The trick that keeps this
    small: register the URL as an external relationship (which writes it into
    word/_rels/document.xml.rels and hands back an id), wrap a fresh run element
    in a <w:hyperlink r:id="..."> element, then hand that run back through
    python-docx's own Run class so the caller can style it with _set_run() just
    like any other run.
    """
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    link.append(r)
    paragraph._p.append(link)            # append() preserves build order
    run = Run(r, paragraph)
    run.text = text
    return run


def _contact_href(field: str, value: str) -> str | None:
    """
    Turn a contact value into a proper link target, or None if it isn't a link.
    The display text stays whatever is in the JSON (e.g. the short
    'linkedin.com/in/you'); only the underlying target gets a scheme added.
    """
    value = value.strip()
    if not value:
        return None
    if field == "email":
        return value if value.startswith("mailto:") else "mailto:" + value
    if field == "phone":
        digits = "".join(c for c in value if c.isdigit() or c == "+")
        return "tel:" + digits if digits else None
    if field in ("website", "linkedin", "github"):
        if value.startswith(("http://", "https://")):
            return value
        return "https://" + value
    return None   # location, etc. -> plain text


# --------------------------------------------------------------------------- #
# Section builders                                                             #
# --------------------------------------------------------------------------- #
def build_header(doc: Document, basics: dict[str, Any], st: Style) -> None:
    """Name, professional title, and a single contact line."""
    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p, after=2)
    _set_run(p.add_run(basics.get("name", "")),
             size=st.name_size, bold=True, color=st.accent, font=st.font_name)

    # Professional title (optional)
    title = basics.get("title", "")
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p, after=4)
        _set_run(p.add_run(title),
                 size=st.title_size, color=st.muted, font=st.font_name)

    # Contact line: only include fields that are present, joined by " | ".
    # Linkable fields (email/phone/website/linkedin/github) become clickable;
    # the display text is unchanged, only the link target gets a scheme.
    fields = ["location", "phone", "email", "website", "linkedin", "github"]
    present = [(f, str(basics[f]).strip()) for f in fields
               if basics.get(f) and str(basics[f]).strip()]
    if present:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p, after=6)
        for i, (field, value) in enumerate(present):
            if i:
                _set_run(p.add_run("  |  "), size=st.contact_size,
                         color=st.muted, font=st.font_name)
            href = _contact_href(field, value)
            run = add_hyperlink(p, value, href) if href else p.add_run(value)
            _set_run(run, size=st.contact_size, color=st.text, font=st.font_name)


def build_section_header(doc: Document, text: str, st: Style) -> None:
    """A bold, uppercase section header with a thin rule beneath it."""
    p = doc.add_paragraph()
    _tight(p, before=8, after=3)
    _set_run(p.add_run(text.upper()),
             size=st.section_size, bold=True, color=st.accent, font=st.font_name)
    _add_bottom_border(p)


def build_summary(doc: Document, summary: str, st: Style) -> None:
    if not summary:
        return
    build_section_header(doc, "Summary", st)
    p = doc.add_paragraph()
    _tight(p, after=2, line=1.05)
    _set_run(p.add_run(summary), size=st.body_size, color=st.text, font=st.font_name)


def _entry_headline(doc: Document, left: str, right: str, st: Style) -> None:
    """Bold left text with a right-aligned date on the same line."""
    p = doc.add_paragraph()
    _tight(p, before=4, after=0)
    _right_tab(p, st.content_width_in)
    _set_run(p.add_run(left),
             size=st.entry_title_size, bold=True, color=st.text, font=st.font_name)
    if right:
        _set_run(p.add_run("\t" + right),
                 size=st.body_size, color=st.muted, font=st.font_name)


def _entry_subline(doc: Document, text: str, st: Style) -> None:
    if not text:
        return
    p = doc.add_paragraph()
    _tight(p, after=1)
    _set_run(p.add_run(text), size=st.body_size, italic=True,
             color=st.muted, font=st.font_name)


def _bullets(doc: Document, items: list[str], st: Style) -> None:
    for item in items or []:
        p = doc.add_paragraph(style="List Bullet")
        _tight(p, after=1, line=1.03)
        _set_run(p.add_run(item), size=st.body_size, color=st.text, font=st.font_name)


def _date_range(entry: dict[str, Any]) -> str:
    start = str(entry.get("start_date", "")).strip()
    end = str(entry.get("end_date", "")).strip()
    if start and end:
        return f"{start} \u2013 {end}"
    return start or end


def build_experience(doc: Document, experience: list[dict], st: Style) -> None:
    if not experience:
        return
    build_section_header(doc, "Experience", st)
    for job in experience:
        position = str(job.get("position", "")).strip()
        company = str(job.get("company", "")).strip()
        headline = " \u2014 ".join(x for x in (position, company) if x)
        _entry_headline(doc, headline, _date_range(job), st)
        _entry_subline(doc, str(job.get("location", "")).strip(), st)
        _bullets(doc, job.get("highlights", []), st)


def build_education(doc: Document, education: list[dict], st: Style) -> None:
    if not education:
        return
    build_section_header(doc, "Education", st)
    for ed in education:
        degree = str(ed.get("degree", "")).strip()
        field = str(ed.get("field", "")).strip()
        left = ", ".join(x for x in (degree, field) if x)
        _entry_headline(doc, left or str(ed.get("institution", "")), _date_range(ed), st)
        # institution + location as the subline
        inst = str(ed.get("institution", "")).strip()
        loc = str(ed.get("location", "")).strip()
        sub = " \u2014 ".join(x for x in (inst, loc) if x) if left else loc
        _entry_subline(doc, sub, st)
        _bullets(doc, ed.get("details", []), st)


def build_certifications(doc: Document, certs: list[dict], st: Style) -> None:
    if not certs:
        return
    build_section_header(doc, "Certifications", st)
    for c in certs:
        name = str(c.get("name", "")).strip()
        issuer = str(c.get("issuer", "")).strip()
        date = str(c.get("date", "")).strip()
        left = name
        if issuer:
            left = f"{name} \u2014 {issuer}" if name else issuer
        _entry_headline(doc, left, date, st)


def _pack_skills(skills: list[str], st: Style) -> list[str]:
    """
    Combine skills into a few readable, comma-separated rows.

    Two modes:
      * Fixed count   : if st.skills_per_row is set, exactly N skills per row.
      * Width-packed  : otherwise greedily pack skills onto a line until adding
                        the next one would exceed ~st.skills_line_width chars.
                        This keeps the rows visually balanced regardless of how
                        long individual skill names are.
    """
    skills = [str(s).strip() for s in skills if str(s).strip()]
    if not skills:
        return []

    if st.skills_per_row and st.skills_per_row > 0:
        rows = [skills[i:i + st.skills_per_row]
                for i in range(0, len(skills), st.skills_per_row)]
        return [", ".join(r) for r in rows]

    rows: list[list[str]] = [[]]
    length = 0
    for skill in skills:
        add = len(skill) + (2 if rows[-1] else 0)  # 2 for ", "
        if rows[-1] and length + add > st.skills_line_width:
            rows.append([skill])
            length = len(skill)
        else:
            rows[-1].append(skill)
            length += add
    return [", ".join(r) for r in rows]


def build_skills(doc: Document, skills: list[str], st: Style) -> None:
    rows = _pack_skills(skills, st)
    if not rows:
        return
    build_section_header(doc, "Skills", st)
    for row in rows:
        p = doc.add_paragraph()
        _tight(p, after=2, line=1.05)
        _set_run(p.add_run(row), size=st.body_size, color=st.text, font=st.font_name)


# --------------------------------------------------------------------------- #
# Document assembly                                                            #
# --------------------------------------------------------------------------- #
def build_resume(data: dict[str, Any], st: Style) -> Document:
    doc = Document()

    # Base document defaults
    normal = doc.styles["Normal"]
    normal.font.name = st.font_name
    normal.font.size = Pt(st.body_size)

    # Page size + margins
    section = doc.sections[0]
    section.page_width = Inches(st.page_width_in)
    section.page_height = Inches(st.page_height_in)
    section.left_margin = section.right_margin = Inches(st.margin_in)
    section.top_margin = section.bottom_margin = Inches(st.margin_in)

    build_header(doc, data.get("basics", {}), st)
    build_summary(doc, data.get("basics", {}).get("summary", ""), st)
    build_experience(doc, data.get("experience", []), st)
    build_education(doc, data.get("education", []), st)
    build_certifications(doc, data.get("certifications", []), st)
    build_skills(doc, data.get("skills", []), st)
    return doc


def load_data(path: Path) -> dict[str, Any]:
    try:
        with path.open(encoding="utf-8") as fh:
            return json.load(fh)
    except FileNotFoundError:
        sys.exit(f"error: input file not found: {path}")
    except json.JSONDecodeError as exc:
        sys.exit(f"error: {path} is not valid JSON: {exc}")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    ap = argparse.ArgumentParser(
        description="Build an ATS-friendly Word resume from a JSON file.")
    ap.add_argument("-i", "--input", required=True, type=Path,
                    help="path to the resume JSON file")
    ap.add_argument("-o", "--output", type=Path, default=Path("resume.docx"),
                    help="output .docx path (default: resume.docx)")
    ap.add_argument("--skills-per-row", type=int, default=None,
                    help="force exactly N skills per row (default: width-packed)")
    ap.add_argument("--skills-line-width", type=int, default=90,
                    help="approx max characters per skills row when width-packing")
    ap.add_argument("--margin", type=float, default=0.6,
                    help="page margin in inches (default: 0.6)")
    return ap.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    st = Style(
        skills_per_row=args.skills_per_row,
        skills_line_width=args.skills_line_width,
        margin_in=args.margin,
    )
    data = load_data(args.input)
    doc = build_resume(data, st)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(args.output))
    print(f"wrote {args.output}")


if __name__ == "__main__":
    main()
